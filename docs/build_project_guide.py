from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = DOCS / "SignaFlow_AI_Project_Guide.docx"
FLOW_IMAGE = DOCS / ".project-guide-workflow.png"
ARCH_IMAGE = DOCS / ".project-guide-architecture.png"
HERO_IMAGE = ROOT / "frontend" / "src" / "assets" / "signaflow-ai-hero.png"

NAVY = "172333"
CORAL = "D94B55"
CYAN = "14969F"
PALE_BLUE = "F1F7F8"
PALE_CORAL = "FDF3F3"
PALE_GRAY = "F5F6F7"
MID_GRAY = "66717E"
LIGHT_BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, name: str = "Arial", size: float | None = None, bold: bool | None = None, color: RGBColor | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LIGHT_BORDER, size: str = "6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_layout(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_header(table.rows[0])
    for row in table.rows:
        prevent_row_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float], font_size: float = 9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    set_repeat_table_layout(table)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        shade_cell(cell, NAVY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        set_run_font(run, size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Inches(widths[col_index])
            shade_cell(cell, PALE_BLUE if row_index % 2 else "FFFFFF")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            run = paragraph.add_run(str(value))
            set_run_font(run, size=font_size, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc: Document, text: str, *, bold_lead: str | None = None, italic: bool = False, align=None):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, size=10.7, bold=True, color=BLACK)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest, size=10.7, color=BLACK)
        rest.italic = italic
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.7, color=BLACK)
        run.italic = italic
    return paragraph


def add_bullets(doc: Document, items: list[str], level: int = 0):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        set_run_font(run, size=10.5, color=BLACK)


def add_numbered(doc: Document, items: list[tuple[str, str]]):
    for title, explanation in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(5)
        run = paragraph.add_run(f"{title}. ")
        set_run_font(run, size=10.6, bold=True, color=BLACK)
        run = paragraph.add_run(explanation)
        set_run_font(run, size=10.6, color=BLACK)


def add_eyebrow(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text.upper())
    set_run_font(run, size=8.5, bold=True, color=RGBColor(217, 75, 85))
    run.font.letter_spacing = Pt(0.5) if hasattr(run.font, "letter_spacing") else None
    return paragraph


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def load_font(size: int, bold: bool = False):
    choices = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
    ]
    for path in choices:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def draw_centered(draw, box, text, font, fill, line_spacing=8):
    x1, y1, x2, y2 = box
    lines = []
    for paragraph in text.split("\n"):
        lines.extend(wrap(paragraph, width=max(12, int((x2 - x1) / (font.size * 0.58)))))
    heights = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        heights.append(bbox[3] - bbox[1])
    total_height = sum(heights) + line_spacing * max(0, len(lines) - 1)
    y = y1 + (y2 - y1 - total_height) / 2
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        width = bbox[2] - bbox[0]
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=font, fill=fill)
        y += height + line_spacing


def draw_arrow(draw, start, end, color=(102, 113, 126), width=7):
    draw.line([start, end], fill=color, width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - 18 * direction, ey - 12), (ex - 18 * direction, ey + 12)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 12, ey - 18 * direction), (ex + 12, ey - 18 * direction)]
    draw.polygon(points, fill=color)


def create_workflow_image():
    img = Image.new("RGB", (1600, 610), "white")
    draw = ImageDraw.Draw(img)
    title_font = load_font(42, True)
    step_font = load_font(29, True)
    small_font = load_font(22, False)
    draw.text((70, 42), "Connected requirement to revenue journey", font=title_font, fill=(23, 35, 51))
    stages = [
        ("1", "Understand", "Natural language"),
        ("2", "Find", "Catalogue search"),
        ("3", "Build", "Complete solution"),
        ("4", "Qualify", "Lead and score"),
        ("5", "Quote", "Database pricing"),
        ("6", "Assist", "Sales follow-up"),
        ("7", "Create", "Campaign draft"),
        ("8", "Measure", "Screen analytics"),
    ]
    box_w, box_h = 315, 150
    xs = [70, 460, 850, 1240]
    ys = [150, 390]
    colors = [(253, 243, 243), (241, 247, 248)]
    for index, (num, title, detail) in enumerate(stages):
        row, col = divmod(index, 4)
        x, y = xs[col], ys[row]
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=24, fill=colors[row], outline=(217, 217, 217), width=3)
        draw.ellipse((x + 20, y + 25, x + 70, y + 75), fill=(217, 75, 85) if row == 0 else (20, 150, 159))
        draw_centered(draw, (x + 20, y + 25, x + 70, y + 75), num, load_font(24, True), "white", 0)
        draw.text((x + 90, y + 24), title, font=step_font, fill=(23, 35, 51))
        draw.text((x + 90, y + 77), detail, font=small_font, fill=(90, 99, 110))
        if col < 3:
            draw_arrow(draw, (x + box_w + 16, y + box_h / 2), (xs[col + 1] - 16, y + box_h / 2))
    draw_arrow(draw, (1397, 315), (1397, 370), color=(20, 150, 159))
    FLOW_IMAGE.parent.mkdir(parents=True, exist_ok=True)
    img.save(FLOW_IMAGE, quality=95)


def create_architecture_image():
    img = Image.new("RGB", (1600, 800), "white")
    draw = ImageDraw.Draw(img)
    title_font = load_font(42, True)
    box_font = load_font(29, True)
    detail_font = load_font(21, False)
    draw.text((70, 42), "How information moves through the platform", font=title_font, fill=(23, 35, 51))
    top = [
        (70, 155, 330, 295, "Customer or sales", "Natural-language request"),
        (455, 155, 715, 295, "React interface", "Secure browser experience"),
        (840, 155, 1100, 295, "FastAPI backend", "Validation and permissions"),
        (1225, 155, 1485, 295, "Workflow services", "AI and business tools"),
    ]
    for x1, y1, x2, y2, title, detail in top:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=(241, 247, 248), outline=(217, 217, 217), width=3)
        draw_centered(draw, (x1 + 12, y1 + 18, x2 - 12, y1 + 78), title, box_font, (23, 35, 51), 3)
        draw_centered(draw, (x1 + 18, y1 + 80, x2 - 18, y2 - 15), detail, detail_font, (90, 99, 110), 3)
    for index in range(3):
        draw_arrow(draw, (top[index][2] + 18, 225), (top[index + 1][0] - 18, 225))
    lower = [
        (130, 475, 510, 675, "Groq", "Requirement extraction\nSales Copilot\nCampaign copy", PALE_CORAL),
        (610, 475, 990, 675, "Local RAG", "Embeddings and catalogue\nSemantic retrieval\nGrounded product IDs", PALE_BLUE),
        (1090, 475, 1470, 675, "Python business rules", "Lead score and permissions\nPricing VAT and totals\nAnalytics and approvals", PALE_GRAY),
    ]
    for x1, y1, x2, y2, title, detail, fill in lower:
        rgb = tuple(int(fill[i:i+2], 16) for i in (0, 2, 4))
        draw.rounded_rectangle((x1, y1, x2, y2), radius=24, fill=rgb, outline=(217, 217, 217), width=3)
        draw_centered(draw, (x1 + 20, y1 + 15, x2 - 20, y1 + 80), title, box_font, (23, 35, 51), 4)
        draw_centered(draw, (x1 + 24, y1 + 82, x2 - 24, y2 - 15), detail, detail_font, (80, 91, 103), 6)
        draw_arrow(draw, ((x1 + x2) / 2, 440), ((x1 + x2) / 2, 470), color=(20, 150, 159))
    draw.line((1355, 295, 1355, 415), fill=(20, 150, 159), width=7)
    draw.line((320, 415, 1280, 415), fill=(20, 150, 159), width=7)
    ARCH_IMAGE.parent.mkdir(parents=True, exist_ok=True)
    img.save(ARCH_IMAGE, quality=95)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.7)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (("Heading 1", 19, 18, 8), ("Heading 2", 13.5, 13, 5), ("Heading 3", 11.5, 9, 3)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[list_name]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("SignaFlow AI Project Guide   |   ")
    set_run_font(run, size=8.5, color=RGBColor(102, 113, 126))
    add_page_number(paragraph)

    doc.core_properties.title = "SignaFlow AI Project Guide"
    doc.core_properties.subject = "Business purpose AI integration workflow and interface guide"
    doc.core_properties.author = "Independent Portfolio Project"


def build_document():
    create_workflow_image()
    create_architecture_image()
    doc = Document()
    configure_document(doc)

    add_eyebrow(doc, "Product and Business Guide")
    title = doc.add_paragraph("SignaFlow AI Project Guide", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph("Business Purpose AI Integration User Journey and Interface Guide")
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.runs[0]
    set_run_font(run, size=16, bold=True, color=BLACK)
    tagline = doc.add_paragraph("From Requirement to Revenue - Powered by AI")
    tagline.paragraph_format.space_after = Pt(18)
    run = tagline.runs[0]
    set_run_font(run, size=11.5, color=RGBColor(20, 150, 159))
    run.italic = True
    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_after = Pt(4)
    image_paragraph.add_run().add_picture(str(HERO_IMAGE), width=Inches(7.0))
    caption = doc.add_paragraph("Illustrative concept showing connected indoor and outdoor digital signage")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(20)
    run = caption.runs[0]
    set_run_font(run, size=8.5, color=RGBColor(102, 113, 126))
    run.italic = True
    add_body(doc, "Prepared for portfolio reviewers and technical evaluators", align=WD_ALIGN_PARAGRAPH.LEFT)
    add_body(doc, "Prototype guide | September 2026", align=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_page_break()

    add_eyebrow(doc, "Executive Overview")
    add_heading(doc, "What SignaFlow AI Does", 1)
    add_body(doc, "SignaFlow AI is an independent sales and advertising platform prototype for a fictional signage business. It helps a customer describe a business need in ordinary language, finds relevant demo products, assembles a complete solution, captures and scores the sales opportunity, prepares a controlled quotation draft, assists the salesperson, generates campaign content, and shows performance analytics.")
    add_body(doc, "The main value is continuity. The requirement, product shortlist, solution, lead, quotation, follow-up, campaign, and analytics remain connected as one journey. A salesperson does not have to rebuild the customer's context at every stage.")
    add_heading(doc, "What This Document Explains", 2)
    add_bullets(doc, [
        "The business problems the platform addresses.",
        "The exact places where Groq, LangGraph, local retrieval, deterministic Python logic, and future computer vision are used.",
        "The kinds of questions the AI can answer and the facts it must never invent.",
        "The complete restaurant demonstration from the first message to campaign analytics.",
        "The purpose, access level, and visible states of every major interface.",
        "The difference between the current prototype, deterministic fallbacks, simulated analytics, and a production deployment.",
    ])
    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(FLOW_IMAGE), width=Inches(7.0))
    caption = doc.add_paragraph("One connected workflow rather than separate AI demonstrations")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(caption.runs[0], size=8.5, color=RGBColor(102, 113, 126))
    caption.runs[0].italic = True
    doc.add_page_break()

    add_eyebrow(doc, "Business Context")
    add_heading(doc, "Problems the Platform Solves", 1)
    add_body(doc, "Signage suppliers often manage broad catalogues across digital signage, LED displays, printing equipment, vinyl, acrylic, ACP, polycarbonate, LED modules, kiosks, and accessories. That range can be a commercial strength, but customers may not know the product names or technical categories needed for their project.")
    add_table(
        doc,
        ["Business problem", "What happens today", "Platform response"],
        [
            ["Customers describe outcomes rather than products", "A restaurant owner may know that menu screens are needed but not which commercial display, player, mount, or CMS to select.", "The Requirement Agent structures the need and the Product Advisor searches by intent."],
            ["Large catalogue takes time to navigate", "Product discovery depends on exact terms, category knowledge, or staff assistance.", "Local semantic retrieval finds products from plain-language use cases."],
            ["Requirements arrive incomplete", "Sales staff repeatedly ask about environment, quantity, location, size, budget, and timeline.", "Progressive clarification asks only the next decision-changing question."],
            ["Solutions can miss supporting items", "A display recommendation may omit playback, content management, mounts, or site checks.", "The Solution Architect adds compatible supporting components and optional upgrades."],
            ["Lead quality is difficult to compare", "A vague enquiry and an urgent quotation request can enter the same queue.", "Python assigns a transparent Cold, Warm, or Hot score from known facts."],
            ["Quotation drafting repeats manual work", "Sales staff must gather product prices and calculate totals and VAT.", "The quotation service reads database prices and performs all calculations in Python."],
            ["Follow-up quality can vary", "Generic messages may ignore the customer's business and project details.", "The Sales Copilot drafts contextual WhatsApp and email messages for approval."],
            ["Digital signage value is hard to measure", "Hardware can be sold without a clear method to evaluate audience response.", "Campaign and anonymous viewing-region analytics create measurable experiments."],
        ],
        [1.65, 2.7, 2.65],
        8.7,
    )
    add_body(doc, "The platform does not claim that the prototype has already increased sales. It provides a working system and a set of future KPIs that a signage business could measure during a controlled pilot.")
    doc.add_page_break()

    add_eyebrow(doc, "System Workflow")
    add_heading(doc, "Complete Customer and Sales Journey", 1)
    add_numbered(doc, [
        ("Customer describes the requirement", "The customer writes what the business is opening, where it is located, what the signage should do, and any known quantity or constraints."),
        ("Requirement is structured", "The system extracts industry, location, use case, indoor or outdoor environment, quantities, size, budget, timeline, installation needs, content needs, preferred product type, and purchase intent."),
        ("One clarification is asked when needed", "The advisor asks only about information that changes the recommendation, such as indoor versus outdoor use or the number of screens."),
        ("Catalogue knowledge is searched", "Local embeddings compare the request with product names, descriptions, use cases, industries, environments, dimensions, materials, specifications, and sources."),
        ("Products are ranked", "Python combines semantic similarity with industry, environment, use case, size, budget, and availability signals."),
        ("A complete solution is assembled", "The system adds primary products, media players, CMS, mounts, compatibility checkpoints, and optional upgrades where relevant."),
        ("A quotation request becomes a lead", "The customer provides contact details, budget, and timeline. The requirement, shortlist, and solution are saved together."),
        ("The lead receives a deterministic score", "The CRM shows a numeric score, Hot, Warm, or Cold temperature, and the exact reasons for the result."),
        ("Sales receives assistance", "The Sales Copilot summarizes the lead, suggests the next action, compares products, proposes cross-sells, and drafts messages."),
        ("A quotation draft is created", "Database prices are used. Missing prices are marked for confirmation. Python calculates line totals, subtotal, VAT, and grand total."),
        ("Campaign content is generated", "The Campaign Agent creates headlines, CTA text, short copy, layout guidance, display duration, schedule guidance, and three creative variants."),
        ("Performance is reviewed", "Sales and campaign dashboards calculate pipeline and screen metrics. Smart Signage shows anonymous ROI and dwell metrics in simulated demo mode."),
    ])
    add_heading(doc, "System Architecture", 2)
    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(ARCH_IMAGE), width=Inches(7.0))
    caption = doc.add_paragraph("The browser never sends the Groq API key and never calls Groq directly")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(caption.runs[0], size=8.5, color=RGBColor(102, 113, 126))
    caption.runs[0].italic = True

    add_eyebrow(doc, "AI Architecture")
    add_heading(doc, "Where AI Is Integrated", 1)
    add_body(doc, "The project deliberately separates language intelligence from authoritative business logic. Groq interprets and drafts language. Local retrieval finds catalogue records. Python controls identities, scores, prices, calculations, permissions, and status changes.")
    add_table(
        doc,
        ["Component", "How intelligence is used", "Authoritative control"],
        [
            ["Requirement Agent", "Groq can return validated structured JSON from a customer's message. A deterministic parser handles the same demo scenarios when Groq is unavailable.", "Pydantic validates every field. Missing facts remain empty and clarification rules run in Python."],
            ["LangGraph workflow", "Maintains journey state and routes from requirement to clarification, recommendation, and solution.", "The current connected /api/chat graph calls bounded application nodes. It does not grant unrestricted tool access."],
            ["Product RAG", "Local embeddings represent catalogue documents and natural-language requests for semantic retrieval.", "Only products already in the database can be returned. Groq is not used to create embeddings."],
            ["Recommendation Engine", "Uses retrieval similarity as one signal and explains matched signals.", "Python adds explicit industry, environment, use-case, size, budget, and availability scores."],
            ["Solution Architect", "Uses the interpreted requirement and ranked products to form a useful package.", "Python chooses database products and quantities and adds compatibility and site-survey checks."],
            ["Sales Copilot", "Groq can summarize a lead, answer a sales question, and draft WhatsApp or email text.", "The application supplies CRM context. Messages remain drafts and are never sent automatically."],
            ["Campaign Agent", "Groq produces structured campaign copy, layout guidance, duration, schedule suggestions, and three variants.", "Pydantic validates the response. Campaigns remain DRAFT and cannot publish automatically."],
            ["Sales and campaign analytics", "The interface presents findings and recommended experiments in business language.", "Database queries and Python calculate every metric. No LLM invents a number."],
            ["Smart Signage", "The production boundary supports YOLO person detection and anonymous tracking. The current demo returns deterministic simulated metrics.", "No facial recognition, identity, biometrics, or sensitive-attribute inference is allowed."],
            ["Visual Advisor", "A future multimodal provider can review a site photo. The current fallback uses the scene description and filename.", "It never estimates final physical dimensions from an uncalibrated image. A manual survey is required."],
        ],
        [1.45, 3.0, 2.55],
        8.4,
    )
    add_heading(doc, "Groq Request Path", 2)
    add_body(doc, "React interface -> FastAPI endpoint -> central GroqLLMService -> configured Groq model -> validated result -> application response.")
    add_bullets(doc, [
        "The Groq API key stays only in the backend .env file.",
        "Model names are configured centrally with GROQ_MODEL, GROQ_FAST_MODEL, and GROQ_REASONING_MODEL.",
        "Timeouts, rate limits, authentication failures, unavailable models, invalid JSON, and schema failures return to a deterministic fallback.",
        "Retrieved web pages, product descriptions, customer text, and database text are treated as untrusted data and cannot override system rules.",
    ])
    doc.add_page_break()

    add_eyebrow(doc, "Answer Capabilities")
    add_heading(doc, "What the AI Can Answer", 1)
    add_table(
        doc,
        ["Area", "Example question", "Expected answer"],
        [
            ["Requirement understanding", "I need screens for my restaurant.", "A structured requirement plus one useful question such as whether the screens are indoor, outdoor, or both."],
            ["Specific product discovery", "I need something outside my shop to show advertisements.", "A shortlist of outdoor-suitable catalogue products with match reasons, source, known specifications, missing-information notices, and confidence."],
            ["Materials discovery", "I need transparent material for window branding.", "Relevant vinyl, acrylic, or polycarbonate categories based on the described use case and environment."],
            ["Complete solution", "Build the full package for four restaurant screens.", "Displays, players, CMS, mounts, quantities, reasons, compatibility checks, and optional upgrades."],
            ["Lead summary", "Summarize this opportunity and tell me what to do next.", "Customer context, score, temperature, status, requirement summary, missing facts, and a recommended next action."],
            ["Lead prioritization", "Which leads should I contact first.", "The highest-scoring CRM leads in priority order with requirement context."],
            ["Product comparison", "Compare commercial display options.", "Environment, size, and database price where available. Unknown price is shown as price to be confirmed."],
            ["Cross-selling", "Suggest relevant upsell opportunities.", "Compatibility-aware checks for mounts, players, CMS, content design, maintenance, and installation scope."],
            ["Sales communication", "Draft a WhatsApp follow-up for this lead.", "A personalized draft referencing the real requirement with a copy button. Nothing is sent automatically."],
            ["Campaign creation", "Create a bilingual weekend restaurant promotion.", "Campaign name, headline, subheadline, CTA, short copy, layout, display duration, schedule guidance, and three variants."],
            ["Performance review", "Which campaign and time window should we test next.", "A recommendation supported by calculated campaign metrics and a confidence score. Schedule changes require approval."],
            ["Visual guidance", "Where could signage fit on this storefront.", "Possible signage positions and product categories, followed by a manual-measurement and site-survey notice."],
        ],
        [1.45, 2.35, 3.2],
        8.5,
    )
    add_heading(doc, "What the AI Must Not Invent", 2)
    add_body(doc, "The system must not invent product names, prices, dimensions, stock, brands, warranty, certifications, discounts, delivery commitments, or installation timelines. If a product fact is missing, the response directs the user to confirm it with the sales team.")
    doc.add_page_break()

    add_eyebrow(doc, "Main Demonstration")
    add_heading(doc, "Restaurant Demo Walkthrough", 1)
    add_body(doc, "The strongest demonstration starts with the exact request below because it contains enough information to produce an immediate recommendation without unnecessary questions.")
    quote = doc.add_paragraph()
    quote.paragraph_format.left_indent = Inches(0.35)
    quote.paragraph_format.right_indent = Inches(0.35)
    quote.paragraph_format.space_before = Pt(5)
    quote.paragraph_format.space_after = Pt(12)
    run = quote.add_run("I am opening a restaurant in Dubai. I need three indoor digital menu screens and one outdoor advertising screen.")
    set_run_font(run, size=12, bold=True, color=BLACK)
    run.italic = True
    add_table(
        doc,
        ["Step", "System action", "What the user sees"],
        [
            ["1 Login", "Validates the account and role and returns a signed access token.", "The role-appropriate navigation and home page."],
            ["2 Open AI Product Advisor", "Loads the advisor with five demonstration scenarios.", "Requirement input, scenario buttons, and a four-stage progress panel."],
            ["3 Submit the restaurant request", "Groq structured extraction runs when configured. The fallback parser is available if the response fails validation.", "Industry Restaurant, location Dubai, use case digital menus plus advertising, environment both, indoor quantity 3, outdoor quantity 1."],
            ["4 Clarification decision", "Python checks whether a missing fact would materially affect the recommendation.", "No clarification for this exact request because environment, quantities, location, industry, and use case are known."],
            ["5 Search products", "RAG retrieves catalogue documents and Python ranks them.", "Approximately three products with match score, reasons, environment, size, price status, and Demo Product or verified-source label."],
            ["6 Build complete solution", "The Solution Architect selects primary displays and adds supporting components.", "Three indoor commercial displays, one weatherproof outdoor display, four media players, CMS, indoor wall mounts, outdoor mounting provision, compatibility checks, and an optional upgrade."],
            ["7 Request quotation", "Opens a controlled contact form and attaches the full requirement, shortlist, and solution.", "Name, company, email, phone, budget, and timeline fields."],
            ["8 Create lead", "Stores the CRM record and calculates the score in Python.", "A success state such as Hot 100 out of 100 when the demonstration fields are complete."],
            ["9 Open the CRM lead", "Loads the saved requirement, contact data, score, budget, timeline, source, and solution.", "Lead detail page with deterministic score ring and Sales Copilot actions."],
            ["10 Draft follow-up", "Groq or the fallback template uses the selected lead as context.", "A requirement-specific WhatsApp or email draft with Copy control."],
            ["11 Prepare quotation", "Reads database prices, calculates totals and VAT, and flags missing prices.", "A DRAFT quotation. The grand total excludes any item whose price requires confirmation."],
            ["12 Generate campaign", "Creates validated copy from business, offer, screen, resolution, audience, language, and tone.", "Responsive portrait, landscape, LED banner, or menu-board preview with three variants."],
            ["13 Review analytics", "Aggregates database campaign metrics and identifies the best campaign and viewing period.", "Impressions, estimated viewers, engagement, hourly trend, and approval-gated experiments marked simulated in demo mode."],
            ["14 Open Smart Signage", "Runs the deterministic demo analysis for a file name, location, screen, duration, and viewing ROI.", "People passed, people in ROI, engagement proxy, average dwell, peak interval, and traffic timeline."],
        ],
        [0.9, 3.05, 3.05],
        8.25,
    )
    add_heading(doc, "Example Sales Follow Up", 2)
    follow_up = add_body(doc, "Hi Ahmed, thank you for sharing details about the three indoor digital menu screens and one outdoor advertising screen for your Dubai restaurant. We have shortlisted a practical solution and can confirm final specifications, availability, and installation requirements with you. Would a quick call today or tomorrow be convenient?", italic=True)
    follow_up.paragraph_format.keep_together = True
    follow_up.paragraph_format.keep_with_next = True
    follow_up_note = add_body(doc, "This is a draft for review and copying. The platform does not send it.")
    follow_up_note.paragraph_format.keep_together = True

    add_eyebrow(doc, "User Experience")
    add_heading(doc, "Interface Guide", 1)
    add_body(doc, "The interface uses role-based navigation. CUSTOMER users receive the advisory journey. SALES users receive CRM, Copilot, quotations, products, and sales dashboards. ADMIN users receive system, ingestion, user, analytics, and opportunity controls in addition to appropriate sales views.")
    add_table(
        doc,
        ["Screen", "Purpose and visible content", "Typical states and access"],
        [
            ["Home", "Introduces Product Advisor, Sales Copilot, Campaign Studio, and Smart Signage and routes the user into the main workflow.", "Ready state. Content varies through role-aware navigation."],
            ["Login", "Accepts email and password and creates the authenticated session.", "Submitting, invalid credentials, success, and redirect states. All roles."],
            ["AI Product Advisor", "Captures a natural-language requirement, shows extracted fields, recommendations, complete solution, quotation modal, and lead success.", "Idle, understanding, clarification, recommending, solution-building, lead-form, success, and error states. All roles."],
            ["Products", "Searches and filters the catalogue and displays description, environment, dimensions, price status, and data-source label.", "Loading, results, no results, and error states. Customer, Sales, and Admin as configured."],
            ["Dashboard", "Shows total leads, Hot leads, pipeline value, quotations, industries, pipeline stages, and calculated insights.", "Calculating, Demo Data or Live Data, charts, and error states. Sales and Admin."],
            ["Leads", "Lists CRM opportunities with requirement, score, temperature, status, budget, and created date. Supports search and temperature filtering.", "Loading, populated table, filtered empty state, and error state. Sales and Admin."],
            ["Lead Details", "Shows the complete qualified requirement, contact details, score explanation, Copilot actions, and quotation creation.", "Loading, lead loaded, Copilot working, draft generated, quotation created, and error states. Sales and Admin."],
            ["Sales Copilot", "Accepts a general sales question or a question grounded in a selected lead. Provides copyable answers and suggested actions.", "No lead, lead selected, ready, typing, answer, and error states. Sales and Admin."],
            ["Quotations", "Displays draft references, customer, date, item count, VAT, confirmed subtotal, missing-price notice, and PDF download.", "No quotes, draft list, PDF download, and error states. Sales and Admin."],
        ],
        [1.2, 3.65, 2.15],
        8.15,
    )
    doc.add_page_break()
    add_heading(doc, "Extended Product Modules", 2)
    add_table(
        doc,
        ["Screen", "Purpose and visible content", "Typical states and access"],
        [
            ["Campaign Studio", "Captures campaign brief fields and renders CSS-based previews across signage formats.", "Empty preview, generating, completed preview, variant selection, export, DRAFT, and error states. Authenticated users."],
            ["Analytics", "Shows campaign impressions, estimated viewers, engagement proxy, peak window, charts, and recommended experiments.", "Calculating, Demo or Simulated label, populated charts, and error states. Sales and Admin."],
            ["Smart Signage", "Accepts a demo video, location, screen, and ROI context and displays anonymous aggregate traffic metrics.", "No file, sample scenario, analyzing, metrics, chart, simulated label, and error states. Sales and Admin."],
            ["Visual Advisor", "Accepts a site photo and scene description and returns observations, categories, possible positions, and measurement safeguard.", "Empty guidance, image preview, reviewing, fallback result, and error states. Authenticated users."],
            ["Website Opportunities", "Maps existing strengths to opportunity, AI enhancement, and business benefit using positive language.", "Loading, opportunity map, and error states. Admin-oriented."],
            ["Business Impact", "Defines baseline and target direction for measurable operating KPIs without claiming results.", "Loading, potential KPI cards, pilot notice, and error states. Admin-oriented."],
            ["Admin", "Lists users and roles and imports validated approved product JSON as clearly labelled Demo Products.", "Loading, validation failure, import success, and permission-restricted states. Admin only."],
            ["Settings", "Shows environment, demo mode, API health, Groq mode, database, RAG backend, and n8n readiness without exposing secrets.", "Healthy or degraded service indicators. Admin only."],
        ],
        [1.2, 3.65, 2.15],
        8.15,
    )
    doc.add_page_break()

    add_eyebrow(doc, "User Experience States")
    add_heading(doc, "How the Interface Communicates System State", 1)
    add_body(doc, "Each state tells the user what the platform is doing, what information is trustworthy, and whether a human decision is still required.")
    add_table(
        doc,
        ["State", "Meaning", "Example"],
        [
            ["Ready", "The page has loaded and accepts input.", "Advisor text area or Copilot question box is available."],
            ["Loading", "The browser is waiting for the backend.", "Understanding requirement, recommending, creating campaign, or calculating metrics."],
            ["Clarification", "The requirement lacks one fact that changes the result.", "Will the signage be used indoors, outdoors, or in both environments."],
            ["Structured requirement", "The customer message has been converted into visible fields.", "Industry, location, use case, environment, and quantities."],
            ["Recommendation", "Grounded catalogue products have been ranked.", "Product card with confidence, matched signals, environment, size, price status, and source label."],
            ["Complete solution", "Primary and supporting products have been assembled.", "Displays, players, CMS, mounts, optional upgrades, and compatibility checkpoints."],
            ["Success", "A requested operation completed.", "Lead created and scored or product records imported."],
            ["Empty", "No records exist or no records match the filter.", "No quotation drafts or no matching products."],
            ["Error", "An API or validation request failed.", "Friendly message without a stack trace or secret information."],
            ["Fallback", "Groq, Chroma, MiniLM, or another optional dependency is unavailable and local logic continues.", "Requirement extraction, Copilot draft, campaign copy, or hashing-based retrieval still works."],
            ["Demo Product", "The record is illustrative and is not presented as verified supplier catalogue data.", "Badge on product and recommendation cards."],
            ["Demo or Simulated Analytics", "The numbers are generated for demonstration and are not measured business outcomes.", "Campaign and Smart Signage dashboards."],
            ["Draft", "A business-critical output awaits review.", "Quotation and campaign remain DRAFT."],
            ["Approval required", "The platform recommends or drafts an action but cannot execute it.", "Messages are copied manually; campaigns and schedules are not published automatically."],
            ["Role restricted", "The logged-in user lacks permission for the route or action.", "Customers cannot open CRM, approve quotes, or access administration."],
        ],
        [1.45, 3.0, 2.55],
        8.45,
    )
    add_heading(doc, "Progressive Advisor State Sequence", 2)
    add_body(doc, "Idle -> Understanding -> Clarification when needed -> Recommending -> Product shortlist -> Solution building -> Complete solution -> Quotation form -> Lead created and scored.")
    doc.add_page_break()

    add_eyebrow(doc, "Authoritative Logic")
    add_heading(doc, "Business Rules and Human Controls", 1)
    add_heading(doc, "Lead Scoring", 2)
    add_body(doc, "Groq never selects the numeric score. Python applies the following transparent rules and caps the result at 100.")
    add_table(
        doc,
        ["Signal", "Points", "Condition"],
        [
            ["Clear requirement", "15", "Requirement contains a meaningful use case and sufficient detail."],
            ["Business identified", "10", "Company or industry is known."],
            ["Quantity provided", "10", "Total, indoor, or outdoor quantity is known."],
            ["Budget provided", "15", "A budget value exists."],
            ["Timeline within 30 days", "20", "Timeline text resolves to 30 days or less or indicates urgency."],
            ["Quotation requested", "20", "Customer explicitly requests a quotation."],
            ["Actionable contact details", "10", "Email and phone or WhatsApp are present."],
        ],
        [2.1, 0.8, 4.1],
        9.0,
    )
    add_body(doc, "Classification: 0 to 39 Cold, 40 to 69 Warm, and 70 to 100 Hot.", bold_lead="Classification:")
    add_heading(doc, "Quotation Rules", 2)
    add_bullets(doc, [
        "Product price must come from the database. Groq cannot create, change, or infer a price.",
        "Python calculates quantity multiplied by unit price, subtotal, VAT, and grand total using decimal arithmetic.",
        "If a price is missing, the item is marked To be confirmed and is excluded from the calculated grand total.",
        "Every generated quotation remains DRAFT and the PDF states Sales Review Required.",
        "Availability, compatibility, installation, delivery, warranty, certification, and discounts require authorized review.",
    ])
    add_heading(doc, "Permissions", 2)
    add_table(
        doc,
        ["Role or agent", "Allowed", "Not allowed"],
        [
            ["Customer", "Use Product Advisor, view recommendations, request quotation, and preview appropriate campaign output.", "Access CRM administration, approve prices, or send campaigns."],
            ["Sales", "View leads, use Copilot, prepare quotation drafts, review products, and use sales dashboards.", "Bypass price controls or publish external communications automatically."],
            ["Admin", "View analytics and settings, review users, and import validated product data.", "Expose secrets or grant AI unrestricted system access."],
            ["Product Advisor Agent", "Search, filter, get, and compare catalogue products.", "Invent products, modify leads, or approve quotations."],
            ["Sales Copilot Agent", "Read approved sales context and draft next actions or messages.", "Delete leads or send messages."],
            ["Quotation Agent", "Create database-backed draft quotations.", "Generate prices, approve, or send official quotations."],
            ["Campaign Agent", "Create draft campaign copy and layout guidance.", "Publish advertisements or change production schedules."],
        ],
        [1.55, 3.0, 2.45],
        8.5,
    )

    add_eyebrow(doc, "Trust and Safety")
    add_heading(doc, "Data Security Hallucination and Privacy Controls", 1)
    add_heading(doc, "Catalogue Grounding", 2)
    add_body(doc, "Every recommendation carries a product ID, source, optional source URL, and retrieval score. The service can return only database products. The current seed catalogue is marked Demo Product and source demo-seed. It must not be described as scraped or verified supplier data.")
    add_heading(doc, "Prompt Injection Protection", 2)
    add_body(doc, "The central LLM instruction treats website text, product descriptions, customer messages, documents, and database content as untrusted data. Retrieved text cannot override system rules, reveal secrets, trigger arbitrary tools, delete data, approve a quotation, or publish a campaign.")
    add_heading(doc, "Credential Protection", 2)
    add_bullets(doc, [
        "GROQ_API_KEY is read only by the backend from the uncommitted .env file.",
        "The frontend does not receive the key and does not call Groq directly.",
        "API responses, logs, database records, screenshots, Dockerfiles, and documentation must not contain the key.",
        "Passwords are stored as salted PBKDF2-SHA256 hashes and access tokens expire.",
    ])
    add_heading(doc, "Computer Vision Privacy", 2)
    add_body(doc, "The future production architecture uses person detection, anonymous tracking IDs, a user-defined screen viewing region, dwell windows, and aggregate counts. It explicitly excludes facial recognition, identity resolution, race, religion, health, political views, biometrics, and other sensitive inference.")
    add_body(doc, "The current Smart Signage page does not analyze the video contents. In demo mode the selected filename and scenario seed deterministic simulated values. The interface labels those values Demo or Simulated Analytics.")
    add_heading(doc, "Visual Advisor Measurement Safety", 2)
    add_body(doc, "The Visual Advisor may suggest likely positions and product categories, but one uncalibrated photograph cannot establish final physical dimensions. Every result states that manual measurement and a site survey are required.")
    doc.add_page_break()

    add_eyebrow(doc, "Deployment Readiness")
    add_heading(doc, "Current Prototype and Production Path", 1)
    add_table(
        doc,
        ["Capability", "Current prototype", "Production path"],
        [
            ["Language understanding", "Groq structured extraction when configured, with deterministic fallback.", "Monitor quality, evaluate prompts, and calibrate against reviewed sales conversations."],
            ["Product knowledge", "Clearly labelled demo catalogue stored in SQLite.", "Import approved supplier product records with sources, ownership, review dates, and compatibility rules."],
            ["Embeddings", "Local 384-dimensional signed hashing always available.", "Local all-MiniLM-L6-v2 embeddings with persistent Chroma or pgvector."],
            ["Database", "SQLite for reliable local demonstration.", "PostgreSQL with migrations, backups, access controls, and retention policy."],
            ["Authentication", "Role-based demo accounts with signed tokens.", "SSO or managed identity, stronger account lifecycle controls, and audit events."],
            ["Sales communications", "Groq or template-based drafts with copy buttons.", "Approved CRM and messaging integration through n8n with human send confirmation."],
            ["Quotations", "Database-backed draft and downloadable PDF.", "Approved price books, quote versioning, commercial authorization, and ERP or CRM integration."],
            ["Campaign preview", "HTML and CSS screen previews and JSON brief export.", "Approved asset workflow, content library, publishing connector, and schedule approval."],
            ["Campaign analytics", "Seeded simulated metrics and calculated recommendations.", "Screen playback telemetry joined to governed anonymous audience aggregates."],
            ["Smart Signage", "Deterministic simulated output. No video-content analysis.", "Separately deployed YOLO and OpenCV edge worker after privacy, legal, signage notice, and retention review."],
            ["Visual Advisor", "Safe scene-description fallback.", "Configurable multimodal provider with explicit accuracy limits and survey handoff."],
            ["Automation", "n8n-ready events, outbound actions disabled.", "Approved workflows for Hot Lead, Quotation Created, and Lead Not Contacted events."],
        ],
        [1.45, 2.65, 2.9],
        8.35,
    )
    add_heading(doc, "Reliable Failure Behavior", 2)
    add_bullets(doc, [
        "Missing or invalid Groq key: requirement, Copilot, and campaign fallbacks continue.",
        "Groq rate limit, timeout, unavailable model, or invalid structured response: safe error metadata is logged and the fallback is used.",
        "MiniLM or Chroma unavailable: local hashing retrieval continues.",
        "Public product ingestion unavailable: cached or approved demo data remains available.",
        "YOLO or OpenCV unavailable: Smart Signage remains a clearly labelled simulation.",
    ])
    doc.add_page_break()

    add_eyebrow(doc, "Business Measurement")
    add_heading(doc, "Business Value and Pilot Measures", 1)
    add_body(doc, "SignaFlow AI should be judged by measurable operating improvements during a pilot, not by unverified revenue claims. The most useful baseline compares the current process with an AI-assisted process for the same product categories and sales team.")
    add_table(
        doc,
        ["Potential KPI", "What to measure", "Expected direction"],
        [
            ["Product search time", "Minutes from requirement received to reviewed shortlist.", "Lower"],
            ["Lead response time", "Time to first useful requirement-specific response.", "Lower"],
            ["Quote turnaround", "Time from qualified need to reviewed draft.", "Lower"],
            ["Lead-to-quote conversion", "Share of qualified leads that receive a quotation.", "Higher"],
            ["Products per quotation", "Relevant primary, supporting, and optional items per reviewed quote.", "Higher when compatible"],
            ["Cross-sell acceptance", "Accepted compatible accessories and services divided by proposed additions.", "Higher"],
            ["Follow-up consistency", "Share of priority leads contacted inside the agreed service level.", "Higher"],
            ["Campaign engagement", "Anonymous ROI presence and dwell during approved campaign tests.", "Higher after controlled testing"],
        ],
        [1.8, 3.7, 1.5],
        8.8,
    )
    add_heading(doc, "Recommended Pilot", 2)
    add_numbered(doc, [
        ("Choose a focused scope", "Select two catalogue categories, one sales team, and a small number of digital signage locations."),
        ("Approve source data", "Import verified product facts, prices, compatibility rules, and source ownership."),
        ("Record a baseline", "Measure search time, response time, quote turnaround, and follow-up consistency for two to four weeks."),
        ("Run side by side", "Compare the AI-assisted workflow with the current process while retaining salesperson approval."),
        ("Review quality weekly", "Check recommendation relevance, missing-information behavior, copied drafts, quote accuracy, and fallback events."),
        ("Decide from evidence", "Expand only when the measured time, quality, conversion, and governance results justify it."),
    ])
    doc.add_page_break()

    add_eyebrow(doc, "Reference")
    add_heading(doc, "Quick Demonstration Reference", 1)
    add_heading(doc, "Recommended Login", 2)
    add_body(doc, "Sales demo account: sales@signaflow.example.com with the local demonstration password demo123. These credentials are for local demonstration only and must not be used in production.")
    add_heading(doc, "Recommended Customer Request", 2)
    add_body(doc, "I am opening a restaurant in Dubai. I need three indoor digital menu screens and one outdoor advertising screen.", italic=True)
    add_heading(doc, "Expected Requirement", 2)
    add_table(
        doc,
        ["Field", "Expected value"],
        [
            ["Industry", "Restaurant"],
            ["Location", "Dubai"],
            ["Environment", "Both indoor and outdoor"],
            ["Indoor quantity", "3"],
            ["Outdoor quantity", "1"],
            ["Use case", "Digital menus and advertising"],
            ["Clarification", "None required for the exact scenario"],
        ],
        [2.2, 4.8],
        9.3,
    )
    add_heading(doc, "Suggested Five Minute Presentation", 2)
    add_numbered(doc, [
        ("Show the Product Advisor", "Submit the restaurant requirement and point out the structured fields and source-labelled recommendations."),
        ("Build the complete solution", "Show how the platform adds players, CMS, mounts, compatibility checks, and an optional upgrade."),
        ("Create the quotation request", "Show the Hot or Warm score and explain that Python, not Groq, calculates it."),
        ("Open the lead", "Generate a WhatsApp draft and prepare the database-priced DRAFT quotation."),
        ("Create a campaign", "Generate a bilingual restaurant promotion and switch between creative variants."),
        ("Show measurement", "Open campaign analytics and Smart Signage and emphasize the simulated labels and privacy boundary."),
    ])
    add_heading(doc, "Main Conclusion", 2)
    add_body(doc, "SignaFlow AI demonstrates a practical way to convert an ordinary customer requirement into a grounded product solution and a qualified sales workflow. AI improves language understanding and drafting, while application code protects product truth, pricing, permissions, calculations, approval, and privacy. The prototype shows how a signage business could make catalogue discovery easier, respond faster, preserve sales context, create more complete solutions, and measure future digital signage services.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
